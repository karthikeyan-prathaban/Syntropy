from __future__ import annotations

import csv
import io
import logging

from app.ingestion.parsers.base import ParseError, clean_text

logger = logging.getLogger(__name__)

Table = list[list[str]]


def is_pdf(data: bytes) -> bool:
    return data[:5] == b"%PDF-"


def is_xls(data: bytes, filename: str = "") -> bool:
    lower = filename.lower()
    return (
        data[:4] == b"PK\x03\x04" and lower.endswith(".xlsx")
    ) or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def decrypt_pdf(data: bytes, password: str | None) -> bytes:
    """Indian banks routinely email password-protected statements (PAN + DOB, etc.)."""
    import pikepdf

    try:
        with pikepdf.open(io.BytesIO(data), password=password or "") as pdf:
            out = io.BytesIO()
            pdf.save(out)
            return out.getvalue()
    except pikepdf.PasswordError as exc:
        raise ParseError(
            "This PDF is password protected. Re-upload with the password."
        ) from exc


def extract_pdf_tables(data: bytes, password: str | None = None) -> tuple[list[Table], str]:
    """Return (tables, full text). Text is used for bank detection and header sniffing."""
    import pdfplumber

    data = decrypt_pdf(data, password)
    tables: list[Table] = []
    text_parts: list[str] = []

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
            for table in page.extract_tables() or []:
                cleaned = [[clean_text(cell) for cell in row] for row in table if any(row)]
                if len(cleaned) >= 2:
                    tables.append(cleaned)

    if not tables:
        # Ruled-line detection fails on some statements; fall back to whitespace columns.
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                found = page.extract_tables(
                    {"vertical_strategy": "text", "horizontal_strategy": "text"}
                )
                for table in found or []:
                    cleaned = [[clean_text(cell) for cell in row] for row in table if any(row)]
                    if len(cleaned) >= 2:
                        tables.append(cleaned)

    return tables, "\n".join(text_parts)


def extract_csv_table(data: bytes) -> Table:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ParseError("Could not decode the CSV file")

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","

    rows = [
        [clean_text(cell) for cell in row]
        for row in csv.reader(io.StringIO(text), delimiter=delimiter)
        if any(str(cell).strip() for cell in row)
    ]
    if not rows:
        raise ParseError("The CSV file is empty")
    return rows


def extract_xls_table(data: bytes, filename: str = "") -> Table:
    if filename.lower().endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        return [
            [clean_text(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)
        ]

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    return [
        [clean_text(cell) for cell in row]
        for row in sheet.iter_rows(values_only=True)
        if any(cell is not None and str(cell).strip() for cell in row)
    ]


def extract(data: bytes, filename: str, password: str | None = None) -> tuple[list[Table], str]:
    """Normalize any supported statement file into tables plus raw text."""
    if is_pdf(data):
        return extract_pdf_tables(data, password)
    if is_xls(data, filename):
        table = extract_xls_table(data, filename)
        return [table], "\n".join(" ".join(row) for row in table[:40])
    table = extract_csv_table(data)
    return [table], "\n".join(" ".join(row) for row in table[:40])
